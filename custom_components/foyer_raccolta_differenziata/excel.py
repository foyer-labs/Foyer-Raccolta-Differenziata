"""Il file Excel della configurazione: scriverlo e leggerlo (decisione 59).

Il formato (fogli, colonne, parole) e le conversioni sono in `core/tabelle.py`; qui
c'è solo quello che riguarda il file: stili, menu a tendina, colonne nascoste, la
guida nel foglio Leggimi, e la lettura delle celle così come le ha salvate Excel,
LibreOffice o Google Fogli.

Il file è organizzato così, in ogni foglio di dati:

    riga 1  titolo
    riga 2  due righe di aiuto
    riga 3  un esempio in grigio
    riga 4  intestazioni (con un commento che spiega la colonna)
    riga 5… i dati

L'esempio sta sopra le intestazioni apposta: la lettura parte dalla riga dopo le
intestazioni, così l'esempio non si importa anche se il foglio di calcolo ne perde
lo stile. Le colonne si riconoscono dal testo dell'intestazione, non dalla
posizione: spostarle o aggiungerne di proprie non rompe niente.

Tutte le funzioni sono sincrone e lente per il loop di Home Assistant: si chiamano
in un executor. openpyxl si importa dentro le funzioni per lo stesso motivo.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any
import warnings
import zipfile

from . import testi
from .core import tabelle as tb

MASSIMO_BYTE = 1024 * 1024
# Un .xlsx è uno zip: un file piccolo può gonfiarsi molto una volta aperto.
MASSIMO_BYTE_ESPANSI = 40 * 1024 * 1024
# Righe vuote di fila dopo le quali la lettura di un foglio si ferma: un foglio
# "formattato fino in fondo" dichiara un milione di righe vuote.
MASSIMO_VUOTE_DI_FILA = 500
RIGA_TITOLO, RIGA_AIUTO, RIGA_ESEMPIO, RIGA_INTESTAZIONI = 1, 2, 3, 4
PRIMA_RIGA = 5
# Le righe preparate con formati e menu, oltre a quelle compilate.
RIGHE_PREPARATE = 200

INCHIOSTRO = "0D1014"
CARTA = "E8ECF2"
AMBRA = "F0A835"
GRIGIO = "8A94A3"
FORMATO_DATA = "DD/MM/YYYY"
FORMATO_ORA = "HH:MM"
FORMATO_TESTO = "@"


class FileNonValido(ValueError):
    """Un file che non si legge come .xlsx; `codice` dice perché."""

    def __init__(self, codice: str) -> None:
        super().__init__(codice)
        self.codice = codice


# --- scrittura ------------------------------------------------------------------------


def _colonna(indice: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(indice)


def _valore_cella(formato: str, valore: Any) -> Any:
    """Il valore da scrivere: orari e date veri, così Excel li mostra come tali."""
    if tb.vuota(valore):
        return None
    if formato == "ora" and isinstance(valore, str):
        return datetime.strptime(tb.leggi_ora(valore), "%H:%M").time()
    return valore


def _formato_numero(formato: str, valore: Any) -> str | None:
    if formato == "data":
        return FORMATO_DATA
    if formato == "ora":
        return FORMATO_ORA
    if formato == "data_o_giorno":
        # Una data vera si mostra come data; "01/06" resta testo, altrimenti Excel
        # lo trasformerebbe nel primo giugno di quest'anno.
        return FORMATO_DATA if isinstance(valore, date) else FORMATO_TESTO
    if formato == "testo":
        return FORMATO_TESTO
    return None


def _contrasto(colore: str) -> str:
    """Testo nero o bianco sopra un colore, per leggibilità."""
    rosso, verde, blu = (int(colore[i : i + 2], 16) for i in (0, 2, 4))
    return "000000" if (0.299 * rosso + 0.587 * verde + 0.114 * blu) > 150 else "FFFFFF"


def _foglio_dati(
    cartella: Any, foglio: tb.Foglio, righe: Sequence[Mapping[str, Any]]
) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = cartella.create_sheet(foglio.nome)
    colonne = foglio.colonne
    # Titolo e aiuto occupano almeno sei colonne: in un foglio stretto (Vacanze)
    # l'aiuto altrimenti si taglia.
    ultima = _colonna(max(len(colonne), 6))
    titolo, aiuto = testi.EXCEL_FOGLI[foglio.nome]

    ws.merge_cells(f"A{RIGA_TITOLO}:{ultima}{RIGA_TITOLO}")
    ws[f"A{RIGA_TITOLO}"] = titolo
    ws[f"A{RIGA_TITOLO}"].font = Font(size=15, bold=True, color=INCHIOSTRO)
    ws.row_dimensions[RIGA_TITOLO].height = 26

    esempio = testi.EXCEL_ESEMPI.get(foglio.nome)
    ws.merge_cells(f"A{RIGA_AIUTO}:{ultima}{RIGA_AIUTO}")
    ws[f"A{RIGA_AIUTO}"] = f"{aiuto} {testi.EXCEL_ESEMPIO}" if esempio else aiuto
    ws[f"A{RIGA_AIUTO}"].font = Font(italic=True, color="4A5563")
    ws[f"A{RIGA_AIUTO}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[RIGA_AIUTO].height = 34

    grigio = Font(italic=True, color=GRIGIO)
    intestazione = Font(bold=True, color=CARTA)
    fondo = PatternFill("solid", fgColor=INCHIOSTRO)
    bordo = Border(bottom=Side(style="medium", color=AMBRA))
    for indice, colonna in enumerate(colonne, start=1):
        lettera = _colonna(indice)
        ws.column_dimensions[lettera].width = colonna.larghezza
        if colonna.nascosta:
            ws.column_dimensions[lettera].hidden = True
        cella = ws.cell(RIGA_INTESTAZIONI, indice, colonna.intestazione)
        cella.font, cella.fill, cella.border = intestazione, fondo, bordo
        cella.alignment = Alignment(vertical="center", wrap_text=True)
        spiegazione = testi.EXCEL_AIUTO_COLONNE.get((foglio.nome, colonna.chiave))
        if spiegazione:
            cella.comment = Comment(spiegazione, "Foyer", width=260, height=90)
        if esempio and colonna.chiave in esempio:
            valore = _valore_cella(colonna.formato, esempio[colonna.chiave])
            cella_esempio = ws.cell(RIGA_ESEMPIO, indice, valore)
            cella_esempio.font = grigio
            if formato := _formato_numero(colonna.formato, valore):
                cella_esempio.number_format = formato
    ws.row_dimensions[RIGA_INTESTAZIONI].height = 30
    ws.freeze_panes = f"A{PRIMA_RIGA}"

    ultima_preparata = PRIMA_RIGA + max(len(righe), RIGHE_PREPARATE) - 1
    for indice, colonna in enumerate(colonne, start=1):
        # Prima i formati di tutte le righe preparate: una data scritta a mano in
        # una riga vuota diventa una data, "1, 15" resta testo.
        formato = _formato_numero(colonna.formato, None)
        if formato is None:
            continue
        for numero in range(PRIMA_RIGA, ultima_preparata + 1):
            ws.cell(numero, indice).number_format = formato
    for numero, riga in enumerate(righe, start=PRIMA_RIGA):
        for indice, colonna in enumerate(colonne, start=1):
            valore = _valore_cella(colonna.formato, riga.get(colonna.chiave))
            cella = ws.cell(numero, indice, valore)
            if formato := _formato_numero(colonna.formato, valore):
                cella.number_format = formato
            if colonna.formato == "data_o_giorno":
                # Date e "gg/mm" nella stessa colonna: allineate allo stesso modo.
                cella.alignment = Alignment(horizontal="left")
            if colonna.chiave == "colore" and isinstance(valore, str):
                cifre = valore.lstrip("#").upper()
                cella.fill = PatternFill("solid", fgColor=cifre)
                cella.font = Font(color=_contrasto(cifre))

    intervallo = PRIMA_RIGA, ultima_preparata
    for indice, colonna in enumerate(colonne, start=1):
        lettera = _colonna(indice)
        celle = f"{lettera}{intervallo[0]}:{lettera}{intervallo[1]}"
        if colonna.scelte:
            menu = DataValidation(
                type="list",
                formula1='"' + ",".join(colonna.scelte) + '"',
                allow_blank=True,
                showErrorMessage=True,
                error=testi.EXCEL_SCELTA_NON_VALIDA,
            )
        elif colonna.chiave == "tipologia":
            nomi = tb.TIPOLOGIE.colonna("nome")
            posizione = _colonna(tb.TIPOLOGIE.colonne.index(nomi) + 1)
            menu = DataValidation(
                type="list",
                formula1=(
                    f"={tb.TIPOLOGIE.nome}!${posizione}${PRIMA_RIGA}:"
                    f"${posizione}${PRIMA_RIGA + 499}"
                ),
                allow_blank=True,
                showErrorMessage=True,
                error=testi.EXCEL_TIPOLOGIA_NON_VALIDA,
            )
        else:
            continue
        ws.add_data_validation(menu)
        menu.add(celle)


def _foglio_impostazioni(cartella: Any, righe: Sequence[Mapping[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    foglio = tb.IMPOSTAZIONI
    ws = cartella.create_sheet(foglio.nome)
    titolo, aiuto = testi.EXCEL_FOGLI[foglio.nome]
    ws.merge_cells(f"A{RIGA_TITOLO}:B{RIGA_TITOLO}")
    ws[f"A{RIGA_TITOLO}"] = titolo
    ws[f"A{RIGA_TITOLO}"].font = Font(size=15, bold=True, color=INCHIOSTRO)
    ws.row_dimensions[RIGA_TITOLO].height = 26
    ws.merge_cells(f"A{RIGA_AIUTO}:B{RIGA_AIUTO}")
    ws[f"A{RIGA_AIUTO}"] = aiuto
    ws[f"A{RIGA_AIUTO}"].font = Font(italic=True, color="4A5563")
    ws[f"A{RIGA_AIUTO}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[RIGA_AIUTO].height = 34
    for indice, colonna in enumerate(foglio.colonne, start=1):
        ws.column_dimensions[_colonna(indice)].width = colonna.larghezza + 8
        cella = ws.cell(RIGA_INTESTAZIONI, indice, colonna.intestazione)
        cella.font = Font(bold=True, color=CARTA)
        cella.fill = PatternFill("solid", fgColor=INCHIOSTRO)
        cella.border = Border(bottom=Side(style="medium", color=AMBRA))
    ws.freeze_panes = f"A{PRIMA_RIGA}"

    voci = {v.etichetta: v for v in tb.VOCI_IMPOSTAZIONI}
    for numero, riga in enumerate(righe, start=PRIMA_RIGA):
        voce = voci[riga["impostazione"]]
        ws.cell(numero, 1, voce.etichetta).font = Font(bold=True)
        valore = _valore_cella(voce.formato, riga.get("valore"))
        cella = ws.cell(numero, 2, valore)
        if voce.formato in ("data", "ora"):
            cella.number_format = _formato_numero(voce.formato, valore) or "General"
        elif voce.formato in ("testo", "giorno_mese"):
            cella.number_format = FORMATO_TESTO
        cella.alignment = Alignment(horizontal="left")
        if voce.scelte:
            menu = DataValidation(
                type="list",
                formula1='"' + ",".join(voce.scelte) + '"',
                allow_blank=True,
                showErrorMessage=True,
                error=testi.EXCEL_SCELTA_NON_VALIDA,
            )
            ws.add_data_validation(menu)
            menu.add(f"B{numero}")


def _foglio_leggimi(
    cartella: Any, destinatari: Sequence[tuple[str, str]], creato: str
) -> None:
    from openpyxl.styles import Alignment, Font

    ws = cartella.active
    ws.title = "Leggimi"
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False
    ws["A1"] = testi.EXCEL_LEGGIMI_TITOLO
    ws["A1"].font = Font(size=16, bold=True, color=INCHIOSTRO)
    ws.row_dimensions[1].height = 30
    numero = 2
    stili = {
        "testo": Font(size=11),
        "sezione": Font(size=13, bold=True, color="B4780F"),
        "punto": Font(size=11),
    }
    for stile, testo in testi.EXCEL_LEGGIMI:
        numero += 2 if stile == "sezione" else 1
        cella = ws.cell(numero, 1, testo)
        cella.font = stili[stile]
        cella.alignment = Alignment(
            wrap_text=True, vertical="top", indent=1 if stile == "punto" else 0
        )
        righe_testo = max(1, len(testo) // 105 + 1)
        ws.row_dimensions[numero].height = 16 * righe_testo + 2
    for identificativo, nome in destinatari or [("", testi.EXCEL_NESSUN_DESTINATARIO)]:
        numero += 1
        testo = f"{identificativo}  —  {nome}" if identificativo else nome
        ws.cell(numero, 1, testo).alignment = Alignment(indent=1)
    numero += 2
    ws.cell(numero, 1, creato).font = Font(italic=True, color=GRIGIO)


def crea(
    righe: Mapping[str, Sequence[Mapping[str, Any]]],
    *,
    destinatari: Sequence[tuple[str, str]],
    creato: str,
) -> bytes:
    """Il file .xlsx: Leggimi e un foglio per sezione, nell'ordine del formato."""
    from openpyxl import Workbook

    cartella = Workbook()
    _foglio_leggimi(cartella, destinatari, creato)
    for foglio in tb.FOGLI:
        if foglio is tb.IMPOSTAZIONI:
            _foglio_impostazioni(cartella, righe.get(foglio.nome, []))
        else:
            _foglio_dati(cartella, foglio, righe.get(foglio.nome, []))
    uscita = BytesIO()
    cartella.save(uscita)
    return uscita.getvalue()


# --- lettura --------------------------------------------------------------------------


def _controlla_zip(contenuto: bytes) -> None:
    if len(contenuto) > MASSIMO_BYTE:
        raise FileNonValido("file_troppo_grande")
    try:
        with zipfile.ZipFile(BytesIO(contenuto)) as archivio:
            espanso = sum(voce.file_size for voce in archivio.infolist())
    except zipfile.BadZipFile as errore:
        raise FileNonValido("file_non_valido") from errore
    if espanso > MASSIMO_BYTE_ESPANSI:
        raise FileNonValido("file_troppo_grande")


def _intestazioni(
    foglio: tb.Foglio, righe: list[tuple[Any, ...]]
) -> dict[int, str] | None:
    """Le colonne della prima riga che ha almeno due intestazioni del formato."""
    for riga in righe:
        trovate: dict[int, str] = {}
        for indice, valore in enumerate(riga):
            chiave = tb.chiave_colonna(foglio, valore)
            if chiave and chiave not in trovate.values():
                trovate[indice] = chiave
        if len(trovate) >= min(2, len(foglio.colonne)):
            return trovate
    return None


def leggi(contenuto: bytes) -> dict[str, tb.Tabella]:
    """I fogli del formato presenti nel file, con le loro righe non vuote."""
    _controlla_zip(contenuto)
    with warnings.catch_warnings():
        # Excel salva i menu che puntano a un altro foglio come estensione, e
        # openpyxl avvisa, mentre legge i fogli, che non la capisce: alla lettura
        # dei valori non serve.
        warnings.simplefilter("ignore", UserWarning)
        return _leggi(contenuto)


def _leggi(contenuto: bytes) -> dict[str, tb.Tabella]:
    from openpyxl import load_workbook

    try:
        cartella = load_workbook(BytesIO(contenuto), read_only=True, data_only=True)
    except Exception as errore:  # openpyxl solleva di tutto su un file rovinato
        raise FileNonValido("file_non_valido") from errore
    try:
        per_nome = {tb.norma(nome): nome for nome in cartella.sheetnames}
        fogli: dict[str, tb.Tabella] = {}
        for foglio in tb.FOGLI:
            nome = per_nome.get(tb.norma(foglio.nome))
            if nome is None:
                continue
            fogli[foglio.nome] = _leggi_foglio(foglio, cartella[nome])
        return fogli
    finally:
        cartella.close()


def _leggi_foglio(foglio: tb.Foglio, ws: Any) -> tb.Tabella:
    righe_iter = ws.iter_rows(values_only=True)
    colonne = None
    numero = 0
    for valori in righe_iter:
        numero += 1
        colonne = _intestazioni(foglio, [valori])
        if colonne is not None or numero >= 20:
            break
    if colonne is None:
        return tb.Tabella(frozenset(), ())
    righe: list[tb.Riga] = []
    vuote = 0
    for valori in righe_iter:
        numero += 1
        celle = {
            chiave: valori[indice] if indice < len(valori) else None
            for indice, chiave in colonne.items()
        }
        if all(tb.vuota(v) for v in celle.values()):
            vuote += 1
            if vuote >= MASSIMO_VUOTE_DI_FILA:
                break
            continue
        vuote = 0
        righe.append(tb.Riga(numero, celle))
        if len(righe) > tb.MASSIMO_RIGHE:
            break  # abbastanza per dire "troppe righe"
    return tb.Tabella(frozenset(colonne.values()), tuple(righe))
