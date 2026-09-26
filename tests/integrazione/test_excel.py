"""La configurazione in Excel dal pannello (decisione 59): scaricare e importare."""

from __future__ import annotations

import base64
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from custom_components.foyer_raccolta_differenziata.const import DOMINIO
from custom_components.foyer_raccolta_differenziata.scambio_excel import URL_EXCEL

from .conftest import configurazione, installa


async def _invia(client, **messaggio):
    await client.send_json_auto_id(messaggio)
    return await client.receive_json()


async def _scarica(hass_client, percorso=URL_EXCEL) -> bytes:
    client = await hass_client()
    risposta = await client.get(percorso)
    assert risposta.status == 200, await risposta.text()
    assert risposta.headers["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return await risposta.read()


def _b64(contenuto: bytes) -> str:
    return base64.b64encode(contenuto).decode()


def _righe(cartella, foglio: str) -> list[tuple]:
    return [
        riga
        for riga in cartella[foglio].iter_rows(min_row=5, values_only=True)
        if any(v is not None for v in riga)
    ]


async def test_il_modello_ha_leggimi_i_fogli_e_le_tipologie_di_base(
    hass, hass_storage, hass_client
):
    await installa(hass, hass_storage)
    contenuto = await _scarica(hass_client, f"{URL_EXCEL}?modello=1")

    cartella = load_workbook(BytesIO(contenuto))
    assert cartella.sheetnames == [
        "Leggimi",
        "Tipologie",
        "Regole",
        "Eccezioni",
        "Promemoria",
        "Vacanze",
        "Impostazioni",
    ]
    tipologie = cartella["Tipologie"]
    assert [r[0] for r in _righe(cartella, "Tipologie")] == [
        "Umido",
        "Carta",
        "Plastica",
        "Vetro",
        "Secco",
        "Verde",
    ]
    # L'intestazione è alla riga 4, l'esempio sopra; l'ID è nascosto.
    assert tipologie["A4"].value == "Nome"
    assert tipologie["A3"].value == "Umido"
    assert tipologie.column_dimensions["H"].hidden
    assert tipologie["A4"].comment is not None
    assert _righe(cartella, "Regole") == []
    menu = {
        str(v.sqref): v.formula1
        for v in cartella["Eccezioni"].data_validations.dataValidation
    }
    assert menu["B5:B204"] == '"Aggiungi,Togli,Sposta"'
    assert menu["A5:A204"].startswith("=Tipologie!$A$5")


async def test_esportare_e_reimportare_non_cambia_nulla(
    hass, hass_storage, hass_client, hass_ws_client
):
    config = configurazione(
        eccezioni=[
            {
                "id": "e1",
                "tipo": "sposta",
                "tipologia": "carta",
                "da": "2026-12-25",
                "a": "2026-12-27",
                "nota": "Natale",
            }
        ],
        patrono={"data": "12-07", "nome": "Sant'Ambrogio"},
        sospensioni=[{"dal": "2026-08-01", "al": "2026-08-20"}],
    )
    await installa(hass, hass_storage, config)
    contenuto = await _scarica(hass_client)
    client = await hass_ws_client(hass)

    for modo in ("sostituisci", "aggiungi"):
        risposta = await _invia(
            client,
            type=f"{DOMINIO}/excel/importa",
            contenuto=_b64(contenuto),
            modo=modo,
        )
        assert risposta["success"], risposta
        esito = risposta["result"]
        assert esito["errori"] == []
        candidata = esito["configurazione"]
        assert candidata["revisione"] == 3
        assert candidata["tipologie"] == config["tipologie"]
        assert candidata["regole"] == config["regole"]
        assert candidata["eccezioni"] == config["eccezioni"]
        assert candidata["patrono"] == config["patrono"]
        assert candidata["sospensioni"] == config["sospensioni"]
        assert all(
            v == {"aggiunte": 0, "modificate": 0, "tolte": 0}
            for v in esito["riepilogo"].values()
        )


async def test_un_file_modificato_in_excel_si_importa_e_si_salva(
    hass, hass_storage, hass_client, hass_ws_client
):
    await installa(hass, hass_storage)
    cartella = load_workbook(BytesIO(await _scarica(hass_client)))
    regole = cartella["Regole"]
    # La prima regola (umido, giovedì) passa anche il lunedì; una riga nuova per il
    # vetro, il 2° e 4° mercoledì, con una tipologia nuova.
    regole["E5"] = "Lun, Gio"
    regole.append(
        ["Vetro", None, "Mensile per giorno della settimana", None, "Mer", "2°, 4°"]
    )
    cartella["Tipologie"].append(["Vetro"])
    uscita = BytesIO()
    cartella.save(uscita)

    client = await hass_ws_client(hass)
    risposta = await _invia(
        client,
        type=f"{DOMINIO}/excel/importa",
        contenuto=_b64(uscita.getvalue()),
        modo="sostituisci",
    )
    esito = risposta["result"]
    assert esito["errori"] == []
    candidata = esito["configurazione"]
    assert candidata["regole"][0]["ricorrenza"]["giorni"] == [0, 3]
    vetro = next(t for t in candidata["tipologie"] if t["nome"] == "Vetro")
    assert vetro["colore"] == "#43a047"
    assert candidata["regole"][2]["tipologia"] == vetro["id"]
    assert esito["riepilogo"]["regole"] == {"aggiunte": 1, "modificate": 1, "tolte": 0}

    salvata = await _invia(
        client,
        type=f"{DOMINIO}/config/salva",
        configurazione=candidata,
        revisione=candidata["revisione"],
    )
    assert salvata["result"]["salvato"] is True
    ritiri = await _invia(
        client, type=f"{DOMINIO}/ritiri", dal="2026-10-12", al="2026-10-14"
    )
    tipologie = {t["id"]: t["nome"] for t in ritiri["result"]["tipologie"]}
    assert [
        (r["data"], tipologie[r["tipologia"]]) for r in ritiri["result"]["ritiri"]
    ] == [
        ("2026-10-12", "Umido"),
        ("2026-10-14", "Vetro"),
    ]


async def test_gli_errori_dicono_dove(hass, hass_storage, hass_ws_client):
    await installa(hass, hass_storage)
    cartella = Workbook()
    ws = cartella.active
    ws.title = "Eccezioni"
    ws.append(["Tipologia", "Cosa", "Data", "Spostato al"])
    ws.append(["Umido", "Togli", date(2026, 12, 25)])
    ws.append(["Plastica", "Aggiungi", "31/12/2026"])
    ws.append(["Carta", "Sposta", "32/12/2026", "02/01/2027"])
    uscita = BytesIO()
    cartella.save(uscita)

    client = await hass_ws_client(hass)
    risposta = await _invia(
        client,
        type=f"{DOMINIO}/excel/importa",
        contenuto=_b64(uscita.getvalue()),
        modo="aggiungi",
    )
    assert risposta["result"]["configurazione"] is None
    assert risposta["result"]["errori"] == [
        {
            "foglio": "Eccezioni",
            "riga": 3,
            "colonna": "Tipologia",
            "codice": "tipologia_sconosciuta",
        },
        {
            "foglio": "Eccezioni",
            "riga": 4,
            "colonna": "Data",
            "codice": "data_non_valida",
        },
    ]


async def test_un_file_che_non_e_excel(hass, hass_storage, hass_ws_client):
    await installa(hass, hass_storage)
    client = await hass_ws_client(hass)
    for contenuto, codice in (
        (_b64(b"non sono un foglio di calcolo"), "file_non_valido"),
        ("%%%", "file_non_valido"),
        ("A" * 2_000_000, "file_troppo_grande"),
    ):
        risposta = await _invia(
            client,
            type=f"{DOMINIO}/excel/importa",
            contenuto=contenuto,
            modo="aggiungi",
        )
        assert risposta["result"]["errori"][0]["codice"] == codice

    vuota = BytesIO()
    Workbook().save(vuota)
    risposta = await _invia(
        client,
        type=f"{DOMINIO}/excel/importa",
        contenuto=_b64(vuota.getvalue()),
        modo="aggiungi",
    )
    assert risposta["result"]["errori"][0]["codice"] == "nessun_foglio"


async def test_solo_gli_amministratori(
    hass, hass_storage, hass_client, hass_ws_client, hass_admin_user
):
    await installa(hass, hass_storage)
    hass_admin_user.groups = []
    client = await hass_client()
    assert (await client.get(URL_EXCEL)).status == 401
    ws = await hass_ws_client(hass)
    risposta = await _invia(
        ws, type=f"{DOMINIO}/excel/importa", contenuto="", modo="aggiungi"
    )
    assert risposta["error"]["code"] == "unauthorized"


async def test_l_indirizzo_firmato_funziona_senza_intestazione(
    hass, hass_storage, hass_ws_client, hass_client_no_auth
):
    """Il pannello apre l'indirizzo firmato: così scarica anche l'app Companion."""
    await installa(hass, hass_storage)
    ws = await hass_ws_client(hass)
    firmato = await _invia(ws, type="auth/sign_path", path=f"{URL_EXCEL}?modello=1")
    client = await hass_client_no_auth()
    assert (await client.get(URL_EXCEL)).status == 401
    risposta = await client.get(firmato["result"]["path"])
    assert risposta.status == 200
    assert (
        "raccolta-differenziata-modello.xlsx" in risposta.headers["Content-Disposition"]
    )


async def test_una_configurazione_non_valida_non_si_esporta(
    hass, hass_storage, hass_client
):
    await installa(hass, hass_storage, configurazione(regole=[{"id": "rotta"}]))
    client = await hass_client()
    assert (await client.get(URL_EXCEL)).status == 409
    assert (await client.get(f"{URL_EXCEL}?modello=1")).status == 200


def test_un_file_salvato_da_excel_si_legge():
    """Un'esportazione aperta, modificata e salvata con Microsoft Excel (it-IT).

    Excel riscrive il file a modo suo: stringhe condivise, menu tra fogli come
    estensione, orari come frazioni di giorno. Il file è nei dati dei test.
    """
    from pathlib import Path

    from custom_components.foyer_raccolta_differenziata import excel
    from custom_components.foyer_raccolta_differenziata.core import tabelle as tb

    contenuto = (Path(__file__).parent / "dati" / "salvato-da-excel.xlsx").read_bytes()
    esito = tb.da_tabelle(
        excel.leggi(contenuto),
        {"revisione": 1},
        "sostituisci",
        genera_id=iter(f"id{i}" for i in range(100)).__next__,
        oggi=date(2026, 9, 26),
    )
    assert esito.errori == []
    config = esito.configurazione
    assert [t["nome"] for t in config["tipologie"]] == [
        "Umido",
        "Carta",
        "Plastica",
        "Vetro",
        "Secco",
    ]
    assert len(config["regole"]) == 7
    assert config["regole"][-1]["periodo"] == {
        "tipo": "annuale",
        "dal": "06-01",
        "al": "09-30",
    }
    assert config["eccezioni"][-1]["data"] == "2026-12-31"
    assert config["esposizione"]["inizio_ora"] == "21:00"
    assert config["patrono"] == {"data": "12-07", "nome": "Sant'Ambrogio"}
